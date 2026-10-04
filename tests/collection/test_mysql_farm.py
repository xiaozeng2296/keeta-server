"""Offline import, request-planning and delivery-coverage regressions."""
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
import zlib
from unittest.mock import Mock,patch

from farm.accounts.importer import parse_curl,prepare_bundle
from farm.storage.mysql import Store,classify,business_day,when
from farm.collection.worker import menu_followups
from farm.delivery.export import merge_menu,coverage,sub_items,write_workbook
from mtgsig.request import refresh_request
import openpyxl


def native_bundle():
    from tests.protocol.test_collection_refresh import identity
    headers={'host':'fooddelivery-eu-3.mykeeta.com','token':'SYNTHETIC','userid':'42',
             'uuid':'OWN-DEVICE','csecuuid':'OWN-DEVICE','appversion':'3.12.500','region':'BR'}
    return {'identity':{k:headers[k] for k in ('token','userid','uuid','csecuuid')},
            'device':identity(),'request':{'method':'POST','headers':headers,
            'url':'https://fooddelivery-eu-3.mykeeta.com/api/accurate/v1/campagin/alita/report?userid=42&uuid=OWN-DEVICE&csecplatform=2&csecpkgname=com.sankuai.sailor.ifooddelivery',
            'body':'{"userCommonParam":{"latitude":"-1","longitude":"-2"}}'},'templates':{}}


class MysqlFarmTests(unittest.TestCase):
    def test_worker_keeps_http_result_without_metering_or_traffic_writes(self):
        from farm.collection.worker import Worker
        bundle,_,_=prepare_bundle(native_bundle())
        bundle['proxy']='http://127.0.0.1:7897'
        store=Mock();store.unseal.return_value=bundle
        worker=Worker(store)
        worker.save_state=Mock();worker.mark_sent=Mock();worker.durable_finish=Mock(return_value={'outcome':'success'})
        worker._release_locks=Mock()
        claim={'account':{},'task':{'id':1,'endpoint':'shopInfo','shop_id':'1','latitude':'-1','longitude':'-2','city_id':'3'},
               'connection':Mock(),'locks':[], 'attempt_id':'synthetic'}
        response=Mock(status_code=200,headers={});response.json.return_value={'code':0,'data':{'shopId':'1','name':'test'}}
        with patch('requests.Session.post',return_value=response) as send:
            result=worker.execute(claim)
        self.assertEqual(result['outcome'],'success')
        self.assertEqual(send.call_args.kwargs['proxies'],{'http':bundle['proxy'],'https':bundle['proxy']})
        self.assertTrue(worker.durable_finish.call_args.kwargs['sent'])
        self.assertNotIn('traffic',worker.durable_finish.call_args.kwargs)
        store.record_traffic.assert_not_called()
        worker._release_locks.assert_called_once()

    def test_imported_home_material_uses_shared_render_b16_rule_in_worker(self):
        from farm.collection.worker import Worker
        from mtgsig.signer import decode_a5,compute_a2
        from farm.storage.mysql import compact
        bundle,_,_=prepare_bundle(native_bundle())
        before='[0,0,0],[1,0,1],[0,0,0],0'
        bundle['device']['base_collect']['b16']=before
        store=Mock();store.unseal.return_value=bundle
        worker=Worker(store);worker.save_state=Mock();worker.mark_sent=Mock()
        worker.durable_finish=Mock(return_value={'outcome':'success'});worker._release_locks=Mock()
        claim={'account':{},'task':{'id':1,'endpoint':'productRender','shop_id':'1','latitude':'-1','longitude':'-2','city_id':'3',
            'payload':compact({'shopCategoryList':[{'shopCategoryId':8,'spuIdList':[9]}]})},
            'connection':Mock(),'locks':[],'attempt_id':'synthetic'}
        response=Mock(status_code=200,headers={});response.json.return_value={'code':0,'data':{'shopCategoryList':[]}}
        def inspect(url,**kwargs):
            mt=json.loads(kwargs['headers']['mtgsig']);a2=mt.pop('a2')
            col=json.loads(decode_a5(mt['a5'],mt['a1'],mt['a3'],mt['a4'])[0])
            self.assertEqual(col['b16'],'[0,0,0],[1,0,1],[1,0,0],0')
            self.assertEqual(a2,compute_a2('POST',url,kwargs['data'].decode(),compact(mt),mt['a1'],
                int(mt['a10'].split(',')[1]),sign_sequence=col['b2']))
            return response
        with patch('requests.Session.post',side_effect=inspect) as send:
            self.assertEqual(worker.execute(claim)['outcome'],'success')
        send.assert_called_once()
        self.assertEqual(bundle['device']['base_collect']['b16'],before)

    def test_delivery_skips_cost_queries_and_excludes_legacy_cost_files(self):
        from farm.collection.executions import export_bundle
        import zipfile
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)
            for name in ('costs.json','costs.xlsx','delivery.xlsx'):(folder/name).write_bytes(b'history')
            store=Mock()
            with patch('farm.collection.executions.export_run',return_value={'items':1}):
                result=export_bundle(store,1,folder)
            self.assertNotIn('costs',result)
            store.rows.assert_not_called()
            self.assertEqual((folder/'costs.json').read_bytes(),b'history')
            with zipfile.ZipFile(result['archive']) as archive:
                self.assertEqual(archive.namelist(),['delivery.xlsx'])

    def test_closed_shop_detail_probe_does_not_reserve_or_send(self):
        from farm.collection.worker import Worker
        for status in (4,'4'):
            with self.subTest(status=status):
                store=Mock()
                store.rows.return_value=[{'response_blob':zlib.compress(json.dumps({'code':0,'data':{'status':status}}).encode())}]
                worker=Worker(store);worker.claim=Mock();worker.execute=Mock()
                with patch('farm.collection.worker.refresh_account_material') as refresh:
                    result=worker.probe(1,'productSpecifics',2)
                self.assertEqual(result['status'],'skipped_closed')
                self.assertFalse(result['sent'])
                refresh.assert_not_called();store.run.assert_not_called()
                store.transaction.assert_not_called();worker.claim.assert_not_called();worker.execute.assert_not_called()

    def test_detail_probe_continues_normal_checks_without_closed_evidence(self):
        from farm.collection.worker import Worker
        for evidence in (None,{'status':3},{'status':None},{}):
            with self.subTest(evidence=evidence):
                store=Mock()
                store.rows.return_value=[] if evidence is None else [{'response_blob':zlib.compress(json.dumps({'data':evidence}).encode())}]
                worker=Worker(store);worker.diagnose=Mock(return_value=[{'reason':'paused'}])
                with patch('farm.collection.worker.refresh_account_material') as refresh:
                    result=worker.probe(1,'productSpecifics',2)
                self.assertEqual(result['status'],'paused')
                refresh.assert_called_once_with(store,1)
                worker.diagnose.assert_called_once_with([1],['productSpecifics'],allow_probe=True)




    def test_event_curl_assembles_schemas_without_foreign_or_invented_fingerprint(self):
        from farm.storage.mysql import PATHS,ENDPOINTS
        original=native_bundle();bundle,status,names=prepare_bundle(original)
        self.assertEqual(status,'ready_captured_incognia');self.assertEqual(set(names),set(ENDPOINTS))
        self.assertEqual(original['templates'],{})
        with tempfile.TemporaryDirectory() as tmp:
            from farm.collection.context import RequestContext
            for name,value in [('identity.json',bundle['identity']),('request.json',bundle['request']),('endpoint_templates.json',bundle['templates'])]:
                (Path(tmp)/name).write_text(json.dumps(value))
            url,headers,body=RequestContext(tmp).build(PATHS['productSpecifics'],'123','-3','-4',spu_id=99)
        self.assertEqual(headers['token'],'SYNTHETIC');self.assertIn('uuid=OWN-DEVICE',url)
        self.assertEqual(json.loads(body)['spuId'],99);self.assertEqual(json.loads(body)['latitude'],'-3')
        self.assertNotIn('fingerPrint',json.loads(body));self.assertEqual(bundle['device']['base_siua'],original['device']['base_siua'])
        self.assertEqual(bundle['account_check_request']['headers']['cookie'],'token=SYNTHETIC')
        self.assertNotIn('token_id',bundle['account_check_request']['url'])
        self.assertEqual(bundle['device']['sign_sequence'],original['device']['sign_sequence'])

    def test_captured_schema_wins_and_unsupported_versions_are_explicit(self):
        from farm.storage.mysql import PATHS
        source=native_bundle();captured=deepcopy(source['request']);captured['url']=captured['url'].replace('/api/accurate/v1/campagin/alita/report',PATHS['shopInfo'])
        captured['body']={'shopId':'1','fingerPrint':'OWN-OBSERVATION'}
        source['templates'][PATHS['shopInfo']]=captured
        prepared,_,_=prepare_bundle(source)
        self.assertEqual(prepared['templates'][PATHS['shopInfo']],captured)
        source=native_bundle();source['request']['headers']['appversion']='unknown'
        prepared,status,names=prepare_bundle(source)
        self.assertEqual(names,[]);self.assertEqual(status,'needs_material')
        self.assertEqual(prepared['material_reasons']['shopInfo'],'unsupported_or_missing_endpoint_schema')

    def test_account_only_is_ready_and_foreign_account_check_is_rejected(self):
        source=native_bundle();source['request']['headers']['appversion']='3.12.401'
        source['request']['url']=source['request']['url'].replace('fooddelivery-eu-3','fooddelivery-eu-9')
        source['request']['headers']['host']='fooddelivery-eu-9.mykeeta.com'
        prepared,status,names=prepare_bundle(source)
        self.assertEqual(status,'ready_captured_incognia');self.assertEqual(names,['accountInfo'])
        for field in ('cookie','token'):
            bad=deepcopy(prepared);bad['account_check_request']['headers'][field]='token=FOREIGN' if field=='cookie' else 'FOREIGN'
            with self.assertRaises(ValueError):prepare_bundle(bad)
        bad=deepcopy(prepared);bad['account_check_request']['url']+='&userid=999'
        with self.assertRaises(ValueError):prepare_bundle(bad)

    def test_curl_parser_never_executes_and_preserves_exact_signed_bytes(self):
        curl='curl -H "Host: fooddelivery-eu-1.mykeeta.com" -H "token: TEST" -H "userid: 123" -H \'mtgsig: {"a0":"2.5"}\' --data-binary \'{"x":"a/b","text":"$(touch /never-execute)"}\' --compressed "https://fooddelivery-eu-1.mykeeta.com/api/v1/shop/productList?userid=123&x=a%2Fb"'
        result=parse_curl(curl)
        self.assertEqual(result['body'],'{"x":"a/b","text":"$(touch /never-execute)"}')
        self.assertTrue(result['url'].endswith('x=a%2Fb'))
        for bad in (curl+'; whoami',curl.replace('userid=123','userid=321'),curl.replace('token: TEST','Cookie: token=DIFFERENT\r\n')):
            with self.assertRaises(ValueError):parse_curl(bad)
        with self.assertRaises(ValueError):parse_curl(curl.replace('--data-binary', '--config'))

    def test_credential_envelope_authenticates_and_hides_plaintext(self):
        store=object.__new__(Store);store.key=lambda:b'K'*32
        payload={'token':'never-in-database-plaintext','counter':8}
        key_id,blob=store.seal(payload)
        self.assertNotIn(payload['token'].encode(),blob)
        self.assertEqual(store.unseal({'credential_blob':blob,'encryption_key_id':key_id}),payload)
        corrupted=blob[:-1]+bytes([blob[-1]^1])
        with self.assertRaises(ValueError):store.unseal({'credential_blob':corrupted,'encryption_key_id':key_id})

    def test_menu_lazy_and_unknown_items_produce_distinct_work(self):
        response={'code':0,'data':{'shopCategoryList':[{'shopCategoryId':7,'spuIdList':[1,2,3,4],
            'spuList':[{'spuId':1,'name':'custom','haveMultiSpecs':1},{'spuId':2,'name':'plain','haveMultiSpecs':0},{'spuId':4,'name':'unknown'}]}]}}
        details,render,missing=menu_followups('productList',response,{})
        self.assertEqual(details,['1','4']);self.assertEqual(render,[{'shopCategoryList':[{'shopCategoryId':7,'spuIdList':[3]}]}]);self.assertFalse(missing)
        result={'data':{'shopCategoryList':[{'shopCategoryId':7,'spuList':[]}]}}
        self.assertEqual(menu_followups('productRender',result,render[0])[2],{'3'})

    def test_merge_render_and_details_preserves_missing_coverage(self):
        menu={'shopCategoryList':[{'shopCategoryId':7,'spuIdList':[1,2],'spuList':[{'spuId':1,'name':'A','haveMultiSpecs':0}]}]}
        merged=merge_menu(menu,[{'data':{'shopCategoryList':[{'spuList':[{'spuId':2,'name':'B','haveMultiSpecs':1}]}]}}])
        check=coverage(merged,{})
        self.assertTrue(check['menu_complete']);self.assertEqual(check['missing_detail_ids'],['2'])
        self.assertEqual(menu['shopCategoryList'][0]['spuList'][0]['spuId'],1)
        check=coverage(merged,{'2':{'data':{'name':'B','skuList':[]}}})
        self.assertTrue(check['custom_details_complete'])

    def test_nested_children_and_numeric_zero_price_are_preserved(self):
        data={'skuList':[{'skuId':1,'groupList':[{'groupId':2,'groupName':'pick','minNumber':1,'maxNumber':2,'hasNestedGroup':1,
            'groupSkuList':[{'groupSkuId':3,'name':'option','originPrice':{'amount':0,'displayText':'0'},'groupList':[{'groupId':4,'groupSkuList':[{'groupSkuId':5,'name':'nested'}]}]}]}]}]}
        sub=sub_items(data);option=sub[0]['groupList'][0]['groupSkuList'][0]
        self.assertEqual(option['originPrice']['amount'],0);self.assertEqual(option['groupList'][0]['groupSkuList'][0]['name'],'nested')
        from farm.delivery.export import CK
        row=CK.spu_to_row({'spuId':1,'minPrice':{'amount':10},'discountPrice':{'amount':0}},'test',1,'1','',0)
        self.assertEqual(row['item_sale_price'],0);self.assertEqual(row['item_discount_price'],0)

    def test_named_variants_survive_without_addon_groups_and_empty_ids_do_not(self):
        data={'haveMultiSpecs':1,'skuList':[
            {'skuId':1,'skuSpec':'Small','originPrice':{'amount':0,'currency':'BRL'},'groupList':[]},
            {'skuId':2,'skuSpec':'Large','originPrice':{'amount':2800,'currency':'BRL'},'groupList':[]}]}
        before=deepcopy(data);result=sub_items(data)
        self.assertEqual([s['skuSpec'] for s in result],['Small','Large'])
        self.assertEqual(result[0]['originPrice']['amount'],0)
        self.assertEqual(result[1]['originPrice']['amount'],2800)
        self.assertTrue(all(s['groupList']==[] for s in result))
        result[0]['originPrice']['amount']=100
        self.assertEqual(data,before)
        self.assertEqual(sub_items({'haveMultiSpecs':1,'skuList':[{'skuId':3,'groupList':[]}]}),[])
        self.assertEqual(sub_items({'haveMultiSpecs':0,'skuList':[{'skuId':4,'skuSpec':'Default','groupList':[]}]}),[])

    def test_closed_and_http200_mismatch_are_not_success(self):
        self.assertEqual(classify('productSpecifics',{'code':201003201,'_http_status':200},product_id='1'),('store_closed',False))
        self.assertEqual(classify('productSpecifics',{'code':0,'data':{'spuId':2,'name':'other'}},product_id='1'),('invalid_payload',False))
        self.assertEqual(classify('productSpecifics',{'code':403,'_http_status':403},product_id='1'),('rejected',False))

    def test_closed_policy_preserves_real_coverage_and_never_waives_main_items(self):
        from farm.storage.mysql import shop_is_closed
        from farm.delivery.export import apply_detail_policy
        self.assertTrue(shop_is_closed({'data':{'status':4}}))
        for response in ({},{'data':{'status':3}},{'data':{'status':99}},{'code':403}):
            self.assertFalse(shop_is_closed(response))
        menu={'shopCategoryList':[{'spuIdList':[1,2],'spuList':[{'spuId':1,'name':'A','haveMultiSpecs':1}]}]}
        check=apply_detail_policy(coverage(menu,{}),True)
        self.assertFalse(check['details_required']);self.assertFalse(check['custom_details_complete'])
        self.assertEqual(check['skipped_detail_ids'],check['missing_detail_ids'])
        self.assertFalse(check['menu_complete']);self.assertEqual(check['missing_main_ids'],['2'])
        self.assertTrue(apply_detail_policy(coverage(menu,{}),False)['details_required'])

    def test_unavailable_policy_preserves_missing_facts_and_requires_named_record(self):
        from farm.delivery.export import apply_detail_policy,unavailable_products
        menu={'shopCategoryList':[{'spuIdList':[1,2,3,4],'spuList':[
            {'spuId':1,'name':'Future sale','haveMultiSpecs':1,'availableStatus':0},
            {'spuId':2,'name':'Unknown','haveMultiSpecs':1},
            {'spuId':3,'availableStatus':0}]}]}
        self.assertEqual(unavailable_products(menu),{'1'})
        check=apply_detail_policy(coverage(menu,{}),False,unavailable_products(menu)|{'3','4'})
        self.assertEqual(check['unavailable_detail_ids'],['1'])
        self.assertIn('1',check['missing_detail_ids'])
        self.assertEqual(check['required_missing_detail_ids'],['2','3','4'])
        self.assertFalse(check['menu_complete']);self.assertFalse(check['details_policy_complete'])
        menu['shopCategoryList'][0].update(spuIdList=[1],spuList=menu['shopCategoryList'][0]['spuList'][:1])
        check=apply_detail_policy(coverage(menu,{}),False,unavailable_products(menu))
        self.assertTrue(check['menu_complete']);self.assertTrue(check['details_policy_complete'])
        self.assertFalse(check['custom_details_complete'])
        self.assertFalse(apply_detail_policy(coverage(menu,{}),False)['details_policy_complete'])

    def test_unavailable_policy_does_not_hide_unrelated_nested_gap(self):
        from farm.delivery.export import apply_detail_policy
        menu={'shopCategoryList':[{'spuList':[
            {'spuId':1,'name':'Unavailable','haveMultiSpecs':1,'availableStatus':0},
            {'spuId':2,'name':'Available','haveMultiSpecs':1,'availableStatus':1}]}]}
        details={'2':{'data':{'name':'Available','hasNestedGroup':1}}}
        check=apply_detail_policy(coverage(menu,details),False,{'1'})
        self.assertFalse(check['details_policy_complete'])
        self.assertEqual(check['required_nested_unresolved_ids'],['2'])

    def test_incognia_compatibility_is_explicit_and_does_not_block_other_refresh(self):
        headers={'region':'BR','token':'T','incog-token':'CAPTURED'}
        url,h,body=refresh_request('https://example.test/?__reqTraceID=OLD',headers,'{}',{'incognia_mode':'captured'},timestamp_ms=1790732176000)
        self.assertEqual(h['incog-token'],'CAPTURED');self.assertNotIn('=OLD',url)
        with self.assertRaises(ValueError):refresh_request(url,headers,body,{},timestamp_ms=1790732176000)

    def test_business_date_and_delivery_columns_match_template(self):
        self.assertEqual(str(business_day(when('2026-09-30T02:16:00Z'))),'2026-09-29')
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'delivery.xlsx'
            write_workbook(path,[],[{'shop_id':'1','item_name':'=not-a-formula','sub_item_json':'[]'}],[],{})
            wb=openpyxl.load_workbook(path)
            self.assertEqual(wb.sheetnames,['店铺信息','菜品信息'])
            self.assertEqual((wb.worksheets[0].max_column,wb.worksheets[1].max_column),(25,18))
            self.assertEqual(wb['菜品信息']['F2'].data_type,'s');wb.close()


if __name__=='__main__':unittest.main()
