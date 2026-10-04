"""Keeta 店铺/菜品采集器(无登录, mtgsig 签名): 读 数据采集任务.xlsx 的 shop_id+采集坐标,
跑 shopInfo/productList, 映射成 交付数据.xlsx 的两张表(店铺信息/菜品信息)。
签名走 RPC(真机 App 自签); 大批量可切离线签(keeta_sign_offline)。"""
import json, sys, time, re, requests, argparse, datetime, os
sys.path.insert(0,".")
import openpyxl

SHOP_COLS=["crawl_lat","crawl_lng","user_type","account_name","shop_id","shop_name","shop_status",
 "is_top_shop","shop_categery","distance","shop_eta","delivery_fee","final_delivery_fee","is_keeta_delivery",
 "shop_coupon_name","shop_coupon_entry_threshold","shop_coupon_discount","maximum_discount_amount","shop_coupon_des",
 "del_coupon_name","del_coupon_entry_threshold","del_coupon_discount","del_coupon_des","create_time","timestamp"]
ITEM_COLS=["shop_id","rank","item_category","item_sug_tag","item_id","item_name","item_describe","item_pic",
 "item_discount","item_original_price","item_sale_price","sub_item_json","create_time","timestamp","has_spec",
 "item_discount_price","item_activitytype","item_activityType_json"]

def g(o,*ks,default=None):
    for k in ks:
        if isinstance(o,dict): o=o.get(k)
        elif isinstance(o,list) and isinstance(k,int) and len(o)>k: o=o[k]
        else: return default
        if o is None: return default
    return o

def nums(s):
    return re.findall(r"\d+", s or "")

def now():
    t=datetime.datetime.now()
    return t.strftime("%Y-%m-%d %H:%M:%S"), int(time.time())

def build_body(path, shop, lat, lng):
    import keeta_client as KC
    fp, headers, body = KC.template(path)
    b=json.loads(body); b["shopId"]=str(shop)
    if isinstance(b.get("location"),dict): b["location"]["latitude"]=str(lat); b["location"]["longitude"]=str(lng)
    if "latitude" in b: b["latitude"]=str(lat)
    if "longitude" in b: b["longitude"]=str(lng)
    return fp, headers, json.dumps(b, ensure_ascii=False)

BID="com.sankuai.sailor.ifooddelivery"
def _ssh(c):
    import subprocess
    subprocess.run(["ssh",os.environ.get("KEETA_DEVICE_SSH_TARGET","keeta-device"),c],
        capture_output=True,text=True)


class ReSigner:
    """自愈签名器: uiopen 前台启动+attach; 会话死了自动重启并重试。"""
    def __init__(self): self.s=None; self._spawn()
    def _spawn(self):
        import keeta_client as KC
        import frida
        try:
            if self.s: self.s.close()
        except Exception: pass
        dev=frida.get_usb_device(timeout=10)
        for a in dev.enumerate_applications():
            if a.identifier==BID and a.pid:
                try: dev.kill(a.pid)
                except Exception: pass
        time.sleep(3)
        _ssh(f"uiopen -b {BID}")
        time.sleep(8)                       # 等前台起来 + SAKGuard 初始化
        self.s=KC.Signer(spawn=False)       # attach 到前台运行的 App(更稳, 非 spawn 态)
    def sign(self, method, url, body):
        for attempt in (1,2):
            try:
                return self.s.sign(method, url, body)
            except Exception as e:
                if attempt==2: raise
                print(f"    [签名器重启] {e}", file=sys.stderr)
                self._spawn()
    def close(self):
        try: self.s.close()
        except Exception: pass

class OfflineReSigner:
    """纯离线签名器: 无设备/无RPC, 用 keeta_K.json 取样身份对任意 (url,body) 算 mtgsig。"""
    def __init__(self, sample):
        from keeta_sign_offline import OfflineSigner
        self.s=OfflineSigner(sample)
    def sign(self, method, url, body):
        mt,_=self.s.sign(url, body); return mt   # method 恒为 POST, 与 canonicalString 一致
    def close(self): pass


IDENT_HDRS=("token","incog-token","uuid","csecuuid","userid","csecuserid")

def _apply_identity(headers, identity):
    if identity:
        for k in IDENT_HDRS:
            if identity.get(k):
                headers[k]=identity[k]
        if identity.get("token"):
            headers["cookie"]=f"token={identity['token']}"

def _resp_json(r):
    """非 JSON(边缘 403/反爬 HTML) 回 {code:HTTP} 而非抛异常, 让上层降级逻辑能识别。"""
    try:
        return r.json()
    except Exception:
        return {"code": r.status_code, "message": (r.text or "")[:120]}

def _post_signed(signer, url, headers, body, proxies=None, session=None):
    headers = dict(headers)
    try:
        # Look up the class protocol: Mock/legacy signers may synthesize any
        # attribute on access and must not accidentally opt into refreshing.
        if callable(getattr(type(signer), "prepare_request", None)):
            url, headers, body = signer.prepare_request(url, headers, body)
        headers["mtgsig"] = signer.sign("POST", url, body)
    finally:
        if hasattr(signer, "persist_counter"):
            # Includes counters reserved before a crypto/signing failure.
            signer.persist_counter()
    sender = session if session is not None else requests
    r = sender.post(url, headers=headers, data=body.encode("utf-8"), timeout=25, proxies=proxies)
    result = _resp_json(r)
    if isinstance(result, dict):
        result["_http_status"] = r.status_code
    return result


def call(signer, path, shop, lat, lng, token=None, proxies=None, identity=None, city=None,
         request_context=None, session=None):
    if request_context is not None:
        if token is not None and token != request_context.identity.get("token"):
            raise ValueError("token override cannot mix accounts with a captured context")
        url, headers, body = request_context.build(path, shop, lat, lng, city=city)
        return _post_signed(signer, url, headers, body, proxies, session)
    fp, headers, body = build_body(path, shop, lat, lng)
    _apply_identity(headers, identity)     # 每账号身份(token/uuid/userid...)覆盖抓包旧值
    if token:                              # 用传入的登录账号 token(覆盖抓包里的旧 token)
        headers["token"]=token
        headers["cookie"]=f"token={token}"
    url=f"https://{KC.HOST}{fp}"
    if city:                               # 每账号城市: 改写 URL 的 ci(签名对新 URL 一致)
        url=re.sub(r'(\bci=)\d+', r'\g<1>'+str(city), url)
    return _post_signed(signer, url, headers, body, proxies, session)

def call_specifics(signer, shop, spuId, lat, lng, token=None, proxies=None, identity=None, city=None,
                   request_context=None, session=None):
    if request_context is not None:
        if token is not None and token != request_context.identity.get("token"):
            raise ValueError("token override cannot mix accounts with a captured context")
        url, headers, body = request_context.build(
            "/api/v1/shop/productSpecifics", shop, lat, lng, city=city, spu_id=spuId)
        return _post_signed(signer, url, headers, body, proxies, session)
    fp, headers, body = KC.template("/api/v1/shop/productSpecifics")
    b=json.loads(body); b["spuId"]=spuId; b["shopId"]=str(shop)
    if lat is not None: b["latitude"]=str(lat)
    if lng is not None: b["longitude"]=str(lng)
    body=json.dumps(b, ensure_ascii=False)
    _apply_identity(headers, identity)
    if token: headers["token"]=token; headers["cookie"]=f"token={token}"
    url=f"https://{KC.HOST}{fp}"
    if city:
        url=re.sub(r'(\bci=)\d+', r'\g<1>'+str(city), url)
    return _post_signed(signer, url, headers, body, proxies, session)

def map_shop(d, shop, clat, clng):
    ct,ts=now()
    di=d.get("deliveryInfo") or {}
    ddi=di.get("deliveryDiscountIntroduction") or {}
    # 配送券(del_coupon): 来自 deliveryDiscountIntroduction / activityShowInfo
    del_name=ddi.get("activityShowInfo") or di.get("deliveryDiscountText") or ""
    del_thr=(nums(di.get("deliveryDiscountText") or del_name) or [""])[0]
    dmatch=re.search(r"R\$?\s*(\d+)\s*off", del_name or "", re.I)
    del_disc=dmatch.group(1) if dmatch else (0 if del_name else "")
    del_des=json.dumps([x for x in (g(ddi,"ruleDetails") or [])], ensure_ascii=False) if ddi.get("ruleDetails") else ""
    # 店铺券(shop_coupon): shopActivityList 里"非配送类"活动才算(配送优惠归 del_coupon)。
    # ponytail: 按交付样本(shop_coupon 多为空)best-effort——排除标题含 deliver 的活动;需更多带店铺券的样本核准规则。
    scname=""
    for act in (d.get("shopActivityList") or []):
        title=(act.get("activityTitle") or "")
        if "deliver" not in title.lower():
            scname=act.get("activityShowInfo") or ""; break
    return {
        "crawl_lat":clat,"crawl_lng":clng,"user_type":-1,"account_name":"ph_非登录","shop_id":str(shop),
        "shop_name":d.get("name"),
        "shop_status":g(d,"openingTimeInfo","title",default=""),
        "is_top_shop":"",
        "shop_categery":d.get("primaryCategoryName"),
        "distance":g(d,"distance","number",default=""),
        "shop_eta":"-".join(nums(di.get("deliveryTimeDesc"))[:2]),
        "delivery_fee":g(di,"deliveryFee","amount",default=""),
        "final_delivery_fee":g(d,"finalDeliveryFee","amount",default=0),
        "is_keeta_delivery":di.get("deliveryDesc") or "",
        "shop_coupon_name":scname,"shop_coupon_entry_threshold":"","shop_coupon_discount":"",
        "maximum_discount_amount":"","shop_coupon_des":"",
        "del_coupon_name":del_name,"del_coupon_entry_threshold":del_thr,"del_coupon_discount":del_disc,
        "del_coupon_des":del_des,"create_time":ct,"timestamp":ts,
    }

def specifics_to_subjson(data):
    """productSpecifics.data.skuList[].groupList -> 交付表 sub_item_json 选项组结构"""
    out=[]
    for sku in (data.get("skuList") or []):
        gl=sku.get("groupList") or []
        if not gl: continue
        groups=[]
        for grp in gl:
            groups.append({
                "groupId":grp.get("groupId"),
                "choices_group_name":grp.get("groupName"),
                "choices_group_min":grp.get("minNumber"),
                "choices_group_max":grp.get("maxNumber"),
                "is_required":"Required" if (grp.get("minNumber") or 0)>0 else "Optional",
                "repeatable":grp.get("repeatable") or 0,
                "hasNestedGroup":grp.get("hasNestedGroup") or 0,
                "groupSkuList":[{
                    "groupSkuId":gs.get("groupSkuId"),"spuId":gs.get("spuId"),
                    "sub_item_name":gs.get("name"),
                    "sub_item_price":(gs.get("originPrice") or {}).get("displayText"),
                    "sub_item_tag":(gs.get("extraInfo") or {}).get("logData","") if gs.get("extraInfo") else "",
                } for gs in (grp.get("groupSkuList") or [])]
            })
        if groups: out.append({"skuId":sku.get("skuId"),"groupList":groups})
    return out

def spu_to_row(spu, cname, rank, shop, ct, ts, sub=""):
    tag=g(spu,"promotionTagInfo","discountTag","text",default="") or g(spu,"productTagInfo","recommendTagForWindow","text",default="")
    acts=spu.get("activityInfoList") or []
    sale_price=g(spu,"discountPrice","amount",default=None)
    if sale_price is None: sale_price=g(spu,"minPrice","amount",default="")
    return {
        "shop_id":str(shop),"rank":rank,"item_category":cname,"item_sug_tag":tag,
        "item_id":spu.get("spuId"),"item_name":spu.get("name"),
        "item_describe":spu.get("description") or "",
        "item_pic":(spu.get("pictures") or [""])[0],
        "item_discount":tag,
        "item_original_price":g(spu,"minPrice","amount",default=""),
        "item_sale_price":sale_price,
        "sub_item_json":sub,"create_time":ct,"timestamp":ts,
        "has_spec":spu.get("haveMultiSpecs") or 0,
        "item_discount_price":sale_price,
        "item_activitytype":g(acts,0,"activityType",default=""),
        "item_activityType_json":json.dumps(acts,ensure_ascii=False) if acts else "",
    }

def map_items(d, shop, signer=None, clat=None, clng=None, token=None, fetch_specs=False, proxies=None,
              details=None):
    """内联菜直接取; 懒加载菜(spuIdList 有、spuList 无)+ 可定制菜 调 productSpecifics 补全 + 取子菜。"""
    ct,ts=now(); rows=[]; rank=0
    for cat in (d.get("shopCategoryList") or []):
        cname=cat.get("shopCategoryName")
        inline={spu.get("spuId"):spu for spu in (cat.get("spuList") or [])}
        ids=cat.get("spuIdList") or list(inline.keys())
        if not ids: ids=list(inline.keys())
        for spuId in ids:
            rank+=1; spu=inline.get(spuId); sub=""
            saved = (details or {}).get(str(spuId))
            if isinstance(saved, dict) and saved.get("code") == 0 and saved.get("data"):
                spu = dict(spu or {}, **saved["data"])
                groups = specifics_to_subjson(saved["data"])
                sub = json.dumps(groups, ensure_ascii=False) if groups else ""
            need_spec = fetch_specs and signer is not None and (spu is None or (spu.get("haveMultiSpecs") in (1,"1")))
            if need_spec:
                try:
                    sp=call_specifics(signer, shop, spuId, clat, clng, token, proxies)
                    if sp and sp.get("code")==0 and sp.get("data"):
                        if spu is None: spu=sp["data"]          # 懒加载菜: 用详情补全
                        sub=json.dumps(specifics_to_subjson(sp["data"]),ensure_ascii=False) if specifics_to_subjson(sp["data"]) else ""
                except Exception: pass
            if spu is None:  # 详情也没取到, 至少留 id
                rows.append({"shop_id":str(shop),"rank":rank,"item_category":cname,"item_id":spuId,
                             "item_name":"","item_describe":"","item_pic":"","item_sug_tag":"","item_discount":"",
                             "item_original_price":"","item_sale_price":"","sub_item_json":"","create_time":ct,
                             "timestamp":ts,"has_spec":"","item_discount_price":"","item_activitytype":"","item_activityType_json":""})
                continue
            rows.append(spu_to_row(spu, cname, rank, shop, ct, ts, sub))
    return rows

def load_tasks(xlsx, limit=None, offset=0):
    wb=openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws=wb.worksheets[0]; rows=list(ws.iter_rows(values_only=True)); wb.close()
    hdr={h:i for i,h in enumerate(rows[0])}
    out=[]
    for r in rows[1+offset:]:                # offset: 跳过已采的前 N 家
        out.append((r[hdr["shop_id"]], r[hdr["crawl_lat"]], r[hdr["crawl_lng"]]))
        if limit and len(out)>=limit: break
    return out

def load_tokens(a):
    toks=[]
    if a.tokens: toks+= [t.strip() for t in a.tokens.split(",") if t.strip()]
    if a.tokens_file:
        for line in open(a.tokens_file, encoding="utf-8"):
            line=line.strip()
            if line and not line.startswith("#"): toks.append(line)
    return toks

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--task", required=True)
    ap.add_argument("--limit", type=int, default=3)
    ap.add_argument("--out", default="exports/legacy-sample.xlsx")
    ap.add_argument("--tokens", help="逗号分隔的登录 token(多账号轮换)")
    ap.add_argument("--tokens-file", help="每行一个 token 的文件")
    ap.add_argument("--save-every", type=int, default=50, help="每 N 家增量存盘")
    ap.add_argument("--delay", type=float, default=2.0, help="每家间隔秒(防风控, 会加抖动)")
    ap.add_argument("--specs", action="store_true", help="调 productSpecifics 补懒加载菜+子菜(完整但请求量大)")
    ap.add_argument("--offset", type=int, default=0, help="跳过任务表前 N 家(接着上次继续)")
    ap.add_argument("--offline", nargs="?", const="/tmp/keeta_K.json", default=None,
                    help="纯离线签名(无设备/无RPC), 值为设备取样 keeta_K.json 路径")
    ap.add_argument("--proxies", help='逗号分隔的代理 URL(如 socks5h://${PROXY_USERNAME}:${PROXY_PASSWORD}@ip:port), 按家轮换换出口IP绕IP风控')
    a=ap.parse_args()
    pxlist=[{"http":u,"https":u} for u in (a.proxies.split(",") if a.proxies else []) if u.strip()] or [None]
    tasks=load_tasks(a.task, a.limit, a.offset)
    if a.offline:
        tokens=[None]                   # 离线=单设备取样身份, 用模板 token(换 token 会与取样设备指纹不符触发风控)
        print(f"[*] 任务 {len(tasks)} 家(offset {a.offset}), 离线签名 sample={a.offline}...", file=sys.stderr)
    else:
        tokens=load_tokens(a) or [None] # None=用模板里的旧 token(仅调试)
        print(f"[*] 任务 {len(tasks)} 家(offset {a.offset}), token {len([t for t in tokens if t])} 个, RPC 采集...", file=sys.stderr)
    wb=openpyxl.Workbook(); ws1=wb.active; ws1.title="店铺信息"; ws1.append(SHOP_COLS)
    ws2=wb.create_sheet("菜品信息"); ws2.append(ITEM_COLS)
    signer=OfflineReSigner(a.offline) if a.offline else ReSigner()
    ok=0; items=0; fail=0
    try:
        for i,(shop,clat,clng) in enumerate(tasks):
            tok=tokens[i % len(tokens)]     # 轮换账号
            px=pxlist[i % len(pxlist)]      # 轮换代理出口IP
            try:
                si=call(signer,"/api/v1/shop/shopInfo",shop,clat,clng,tok,px)
                if si.get("code")!=0:
                    fail+=1; print(f"  ! {shop} shopInfo code={si.get('code')} {si.get('message')}",file=sys.stderr); continue
                srow=map_shop(si["data"], shop, clat, clng)
                ws1.append([srow[c] for c in SHOP_COLS])
                pl=call(signer,"/api/v1/shop/productList",shop,clat,clng,tok,px)
                if pl.get("code")==0:
                    for it in map_items(pl["data"], shop, signer, clat, clng, tok, fetch_specs=a.specs, proxies=px):
                        ws2.append([it[c] for c in ITEM_COLS]); items+=1
                ok+=1
                if ok%20==0: print(f"  ..{ok}/{len(tasks)} 家, 菜品 {items}",file=sys.stderr)
                if ok%a.save_every==0: wb.save(a.out)   # 增量存盘防中断
                import random; time.sleep(a.delay + random.uniform(0, a.delay))  # 抖动降风控
            except Exception as e:
                fail+=1; print(f"  ! shop {shop} 异常: {e}",file=sys.stderr)
    finally:
        signer.close()
    wb.save(a.out)
    print(f"\n[done] 成功 {ok} 家, 失败 {fail}, 菜品 {items} 条 -> {a.out}")

if __name__=="__main__": main()
