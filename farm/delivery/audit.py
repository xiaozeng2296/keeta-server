"""Verify the delivery workbook, embedded long fields, coverage, and ZIP."""
import argparse
import hashlib
import json
from pathlib import Path
import zipfile
import openpyxl


def audit(output, expected_shops=None):
    output=Path(output);errors=[];xlsx=output/'delivery.xlsx'
    coverage=json.loads((output/'delivery.coverage.json').read_text())
    checks=coverage['shops'];expected={str(r['shop_id']) for r in checks}
    if expected_shops is not None and expected!=set(map(str,expected_shops)):errors.append('shop_scope_mismatch')
    wb=openpyxl.load_workbook(xlsx,read_only=True,data_only=False)
    counts=[];references=[]
    try:
        if wb.sheetnames[:2]!=['店铺信息','菜品信息'] or set(wb.sheetnames)-{'店铺信息','菜品信息','长字段内容'}:errors.append('sheet_schema_mismatch')
        for name in ('店铺信息','菜品信息'):
            rows=wb[name].iter_rows();keys=[c.value for c in next(rows)];count=0;shopids=set()
            for row in rows:
                count+=1;value=dict(zip(keys,[c.value for c in row]));shopids.add(str(value['shop_id']))
                if any(c.data_type=='f' for c in row):errors.append('unexpected_formula')
                for col,cell in zip(keys,row):
                    if isinstance(cell.value,str) and cell.value.startswith('{'):
                        try:pointer=json.loads(cell.value)
                        except ValueError:continue
                        if pointer.get('sheet')=='长字段内容':references.append(pointer)
            if name=='店铺信息' and shopids!=expected:errors.append('workbook_shop_rows_missing')
            if not shopids<=expected:errors.append('unexpected_workbook_shop')
            counts.append(count)
        if counts!=[coverage['shop_rows'],coverage['item_rows']]:errors.append('row_counts_mismatch')
        fragments={}
        if '长字段内容' in wb:
            for row in wb['长字段内容'].iter_rows(min_row=2):
                if any(c.data_type=='f' for c in row):errors.append('unexpected_formula')
                values=[c.value for c in row];fragments.setdefault(values[0],[]).append(values)
        for pointer in references:
            chunks=fragments.get(pointer['reference'],[])
            if len(chunks)!=pointer['chunks'] or [r[4] for r in chunks]!=list(range(1,pointer['chunks']+1)):
                errors.append('long_field_chunks_missing');continue
            raw=''.join(r[6] or '' for r in chunks).encode()
            if hashlib.sha256(raw).hexdigest()!=pointer['sha256']:errors.append('long_field_hash_mismatch')
        if len(references)!=len(coverage.get('overflow',[])):errors.append('long_field_reference_count')
    finally:wb.close()
    archive=output/'delivery.zip'
    with zipfile.ZipFile(archive) as z:
        if z.namelist()!=['delivery.xlsx']:errors.append('unexpected_zip_members')
        elif hashlib.sha256(z.read('delivery.xlsx')).digest()!=hashlib.sha256(xlsx.read_bytes()).digest():errors.append('zip_workbook_mismatch')
    if any(not r.get('complete') for r in checks):errors.append('incomplete_coverage')
    return dict(ok=not errors,errors=sorted(set(errors)),shops=len(expected),items=coverage['item_rows'],long_fields=len(references))


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('output',type=Path);a=p.parse_args()
    result=audit(a.output);print(json.dumps(result,ensure_ascii=False));return int(not result['ok'])

if __name__=='__main__':raise SystemExit(main())
