"""Export database results using exactly the delivery workbook's two schemas."""
from copy import deepcopy
from datetime import timezone
import json
import hashlib
import os
import tempfile
from pathlib import Path
from farm.storage.responses import load_response, save_response
import openpyxl
from openpyxl.styles import Font,PatternFill

from farm.delivery import mapping as CK
from farm.storage.mysql import compact,digest,unpack_json,closed_shop_jobs


def merge_menu(menu,render_responses):
    menu=deepcopy(menu);categories=menu.get('shopCategoryList') or []
    rendered={}
    for response in render_responses:
        for cat in response.get('data',{}).get('shopCategoryList') or []:
            for product in cat.get('spuList') or []:
                if product.get('spuId') is not None:rendered[str(product['spuId'])]=product
    for cat in categories:
        products={str(p['spuId']):p for p in cat.get('spuList') or [] if p.get('spuId') is not None}
        ids=list(dict.fromkeys([str(x) for x in cat.get('spuIdList') or []]+list(products)))
        for identifier in ids:
            if identifier in rendered:products[identifier]=rendered[identifier]
        cat['spuIdList']=ids
        cat['spuList']=[dict(p,spuId=str(p['spuId'])) for p in products.values()]
    return menu


def sub_items(data):
    """Keep delivered aliases and the original recursive structures together."""
    groups=CK.specifics_to_subjson(data)
    source={str(s.get('skuId')):s for s in data.get('skuList') or []}
    present={str(s['skuId']) for s in groups}
    # A named variant is real specification data even without add-on groups.
    # Bare SKU IDs from an unexpanded menu are still not specification data.
    for identifier,sku in source.items():
        if identifier not in present and sku.get('skuSpec') and (len(source)>1 or data.get('haveMultiSpecs') in (1,'1')):
            groups.append({'skuId':sku.get('skuId'),'groupList':[]})
    for sku in groups:
        original=source.get(str(sku['skuId']),{})
        for name in ('skuSpec','originPrice','discountPrice','activityInfoList','discountPriceList','finalPrice','couponInfo','combinationSpec','combinationSpecList'):
            if name in original:sku[name]=deepcopy(original[name])
        by_group={str(g.get('groupId')):g for g in original.get('groupList') or []}
        for group in sku['groupList']:
            raw=by_group.get(str(group['groupId']),{})
            by_option={str(o.get('groupSkuId')):o for o in raw.get('groupSkuList') or []}
            for option in group['groupSkuList']:
                native=by_option.get(str(option['groupSkuId']),{})
                # Preserve nested structure and the numeric/currency price.
                for name in ('groupList','skuList','originPrice','minNumber','maxNumber','repeatable'):
                    if name in native:option[name]=deepcopy(native[name])
            if raw.get('groupList'):group['groupList']=deepcopy(raw['groupList'])
    order={identifier:i for i,identifier in enumerate(source)}
    groups.sort(key=lambda sku:order.get(str(sku['skuId']),len(order)))
    return groups


def unresolved_nested(data):
    if isinstance(data,list):return any(unresolved_nested(x) for x in data)
    if not isinstance(data,dict):return False
    if data.get('hasNestedGroup') not in (None,False,0,'0'):
        if not data.get('groupList') and not any(x.get('groupList') for x in data.get('groupSkuList') or []):return True
    return any(unresolved_nested(v) for v in data.values() if isinstance(v,(list,dict)))


def unavailable_products(menu):
    """Only explicit menu evidence; absent/unknown availability is not a skip."""
    products={str(p['spuId']):p for c in menu.get('shopCategoryList') or []
              for p in c.get('spuList') or [] if p.get('spuId') is not None}
    return {identifier for identifier,p in products.items()
            if p.get('name') and type(p.get('availableStatus')) in (int,str)
            and p['availableStatus'] in (0,'0')}


def coverage(menu,details):
    ids=set();products={}
    for cat in menu.get('shopCategoryList') or []:
        ids.update(str(x) for x in cat.get('spuIdList') or [])
        for product in cat.get('spuList') or []:
            identifier=str(product['spuId']);ids.add(identifier);products[identifier]=product
    missing_main=[];missing_details=[];nested=[]
    for identifier in sorted(ids):
        detail=details.get(identifier,{}).get('data') or {};product=products.get(identifier) or detail
        if not product.get('name'):missing_main.append(identifier)
        flag=product.get('haveMultiSpecs')
        if flag not in (0,'0') and not detail.get('name'):missing_details.append(identifier)
        if detail and unresolved_nested(detail):nested.append(identifier)
    return {'declared_products':len(ids),'missing_main_ids':missing_main,'missing_detail_ids':missing_details,
            'recorded_product_ids':sorted(ids-set(missing_main)),
            'unavailable_product_ids':sorted(unavailable_products(menu)),
            'nested_unresolved_ids':nested,'menu_complete':not missing_main,
            'custom_details_complete':not missing_details and not nested}


def apply_detail_policy(check,closed,unavailable_ids=()):
    # Keep raw missing/full-data coverage unchanged. Only the policy gate may
    # exempt named, recorded products, never missing menu rows or unknown IDs.
    exempt=set(unavailable_ids)&set(check.get('recorded_product_ids',[]))
    unresolved=set(check['missing_detail_ids'])|set(check['nested_unresolved_ids'])
    check['details_required']=not closed
    check['details_skip_reason']='store_closed' if closed else None
    check['unavailable_detail_ids']=sorted(unresolved&exempt) if not closed else []
    check['skipped_detail_ids']=sorted(unresolved if closed else unresolved&exempt)
    check['required_missing_detail_ids']=sorted(set(check['missing_detail_ids'])-exempt) if not closed else []
    check['required_nested_unresolved_ids']=sorted(set(check['nested_unresolved_ids'])-exempt) if not closed else []
    check['details_policy_complete']=bool(closed or not (unresolved-exempt))
    return check


def text_chunks(value, limit=30000):
    """Excel counts UTF-16 units; never split a non-BMP character in half."""
    start=0;units=0
    for index,char in enumerate(value):
        width=2 if ord(char)>0xffff else 1
        if units+width>limit:
            yield value[start:index];start=index;units=0
        units+=width
    yield value[start:]


def write_workbook(path,shop_rows,item_rows,coverage_rows,source_metadata):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    wb=openpyxl.Workbook();ws=wb.active;ws.title='店铺信息'
    overflow=[];long_sheet=None
    for sheet,columns,rows in ((ws,CK.SHOP_COLS,shop_rows),(wb.create_sheet('菜品信息'),CK.ITEM_COLS,item_rows)):
        sheet.append(columns);sheet.freeze_panes='A2';sheet.auto_filter.ref=f'A1:{openpyxl.utils.get_column_letter(len(columns))}1'
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='334155')
        for row_number,row in enumerate(rows,2):
            values=[]
            for key in columns:
                value=row.get(key,'');value='' if value is None else value
                if isinstance(value,str) and len(value.encode('utf-16-le'))//2>32767:
                    if long_sheet is None:
                        long_sheet=wb.create_sheet('长字段内容')
                        long_sheet.append(['引用ID','来源工作表','来源行','字段','分片序号','分片总数','内容','SHA256'])
                        long_sheet.freeze_panes='A2'
                    sha=hashlib.sha256(value.encode()).hexdigest()
                    ref=f'{sheet.title}:{row_number}:{key}'
                    chunks=list(text_chunks(value));start_row=long_sheet.max_row+1
                    for index,chunk in enumerate(chunks,1):
                        long_sheet.append([ref,sheet.title,row_number,key,index,len(chunks),chunk,sha])
                        for cell in long_sheet[long_sheet.max_row]:
                            if isinstance(cell.value,str):cell.data_type='s'
                    pointer=dict(sheet=long_sheet.title,reference=ref,first_row=start_row,
                                 chunks=len(chunks),sha256=sha,truncated=False)
                    overflow.append(dict(pointer,shop_id=row.get('shop_id'),item_id=row.get('item_id'),column=key))
                    value=compact(pointer)
                values.append(value)
            sheet.append(values)
            for cell in next(sheet.iter_rows(min_row=row_number,max_row=row_number,min_col=1,max_col=len(columns))):
                if isinstance(cell.value,str):cell.data_type='s'
    fd,temporary=tempfile.mkstemp(dir=path.parent,prefix='.workbook-',suffix='.xlsx');os.close(fd)
    try:
        wb.save(temporary);os.replace(temporary,path)
    finally:
        if os.path.exists(temporary):os.unlink(temporary)
    from farm.storage.files import atomic_json
    metadata=dict(source_metadata,shop_rows=len(shop_rows),item_rows=len(item_rows),shops=coverage_rows,overflow=overflow)
    atomic_json(path.with_suffix('.coverage.json'),metadata)
    return {'path':str(path.resolve()),'shops':len(shop_rows),'items':len(item_rows),
            'complete_shop_jobs':sum(r['complete'] for r in coverage_rows),'partial_shop_jobs':sum(not r['complete'] for r in coverage_rows),'overflow_cells':len(overflow)}


def export_run(store,run_id,path):
    run=store.rows('SELECT settings FROM collection_runs WHERE id=%s',(run_id,))
    settings=unpack_json(run[0]['settings']) if run else {}
    closed=closed_shop_jobs(store,run_id) if settings.get('skip_closed_details') is True else set()
    jobs=store.rows('SELECT * FROM shop_jobs WHERE run_id=%s ORDER BY id',(run_id,))
    results=store.rows('SELECT r.*,a.user_id FROM task_results r JOIN shop_jobs j ON j.id=r.shop_job_id LEFT JOIN accounts a ON a.id=r.account_id WHERE j.run_id=%s ORDER BY r.observed_at,r.id',(run_id,))
    grouped={}
    for row in results:
        row['response']=load_response(row['response_blob']);grouped.setdefault(row['shop_job_id'],[]).append(row)
    shops=[];items=[];checks=[]
    for job in jobs:
        raw=grouped.get(job['id'],[]);infos=[r for r in raw if r['endpoint']=='shopInfo'];menus=[r for r in raw if r['endpoint']=='productList']
        info=infos[-1] if infos else None;menu=menus[-1] if menus else None
        if info:
            row=CK.map_shop(info['response']['data'],job['shop_id'],job['latitude'],job['longitude'])
            row.update(user_type=1,account_name=info['user_id'] or '',create_time=info['observed_at'].isoformat(sep=' '),timestamp=int(info['observed_at'].replace(tzinfo=timezone.utc).timestamp()));shops.append(row)
        if menu:
            renders=[r['response'] for r in raw if r['endpoint']=='productRender' and r['observed_at']>=menu['observed_at']]
            merged=merge_menu(menu['response']['data'],renders)
            details={r['target_id']:r['response'] for r in raw if r['endpoint']=='productSpecifics'}
            check=coverage(merged,details)
            rows=CK.map_items(merged,job['shop_id'],details=details)
            for row in rows:
                detail=details.get(str(row['item_id']))
                if detail:
                    sub=sub_items(detail['data']);row['sub_item_json']=compact(sub) if sub else ''
                row.update(create_time=menu['observed_at'].isoformat(sep=' '),timestamp=int(menu['observed_at'].replace(tzinfo=timezone.utc).timestamp()))
            items.extend(rows)
        else:check={'menu_complete':False,'custom_details_complete':False,'declared_products':None,'missing_main_ids':[],'missing_detail_ids':[],'nested_unresolved_ids':[]}
        unavailable=set()
        if settings.get('skip_unavailable_details') is True and menu:
            unavailable=unavailable_products(merged)
            unavailable.update(str(r['target_id']) for r in store.rows(
                "SELECT target_id FROM tasks WHERE shop_job_id=%s AND endpoint='productSpecifics' AND state='skipped_unavailable'",(job['id'],)))
        apply_detail_policy(check,job['id'] in closed,unavailable)
        full=bool(info) and bool(menu) and check['menu_complete'] and check['custom_details_complete']
        check.update(shop_job_id=job['id'],shop_id=job['shop_id'],shop_info_complete=bool(info),
                     full_data_complete=full,
                     complete=bool(info) and bool(menu) and check['menu_complete'] and check['details_policy_complete'])
        checks.append(check)
    return write_workbook(path,shops,items,checks,{'run_id':run_id,'source':'mysql','timestamps':'UTC','atomic_cross_account_snapshot':False,'skip_closed_details':settings.get('skip_closed_details') is True})


def export_status(store,path):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    tables=[('账号状态',store.rows('SELECT * FROM account_dashboard ORDER BY id')),
            ('每日接口统计',store.rows('SELECT * FROM daily_account_requests ORDER BY business_date,account_id,endpoint,origin')),
            ('每日预算',store.rows('SELECT a.label,p.endpoint,p.work_limit,p.hard_limit,p.timezone FROM budget_policies p JOIN accounts a ON a.id=p.account_id ORDER BY a.id,p.endpoint')),
            ('队列状态',store.rows('SELECT j.run_id,t.endpoint,t.state,COUNT(*) AS task_count FROM tasks t JOIN shop_jobs j ON j.id=t.shop_job_id GROUP BY j.run_id,t.endpoint,t.state'))]
    wb=openpyxl.Workbook();wb.remove(wb.active)
    for name,rows in tables:
        ws=wb.create_sheet(name)
        if not rows:continue
        keys=list(rows[0]);ws.append(keys);ws.freeze_panes='A2'
        for row in rows:
            ws.append([row[k] for k in keys])
    wb.save(path)
    return {'path':str(path.resolve()),'sheets':[x[0] for x in tables]}
