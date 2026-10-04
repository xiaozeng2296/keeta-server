"""Export database results using exactly the delivery workbook's two schemas."""
from copy import deepcopy
from datetime import timezone
import json
from pathlib import Path
import zlib
import openpyxl
from openpyxl.styles import Font,PatternFill

from farm import _paths
import crawl_keeta as CK
from farm.mysql_store import compact,digest,unpack_json,closed_shop_jobs


def merge_menu(menu,render_responses):
    menu=deepcopy(menu);categories=menu.get('shopCategoryList') or []
    rendered={}
    for response in render_responses:
        for cat in response.get('data',{}).get('shopCategoryList') or []:
            for product in cat.get('spuList') or []:
                if product.get('spuId') is not None:rendered[str(product['spuId'])]=product
    for cat in categories:
        products={str(p['spuId']):p for p in cat.get('spuList') or [] if p.get('spuId') is not None}
        ids=list(dict.fromkeys([str(x) for x in cat.get('spuIdList') or []]+list(products)))
        for identifier in ids:
            if identifier in rendered:products[identifier]=rendered[identifier]
        cat['spuIdList']=ids
        cat['spuList']=[dict(p,spuId=str(p['spuId'])) for p in products.values()]
    return menu


def sub_items(data):
    """Keep delivered aliases and the original recursive structures together."""
    groups=CK.specifics_to_subjson(data)
    source={str(s.get('skuId')):s for s in data.get('skuList') or []}
    present={str(s['skuId']) for s in groups}
    # A named variant is real specification data even without add-on groups.
    # Bare SKU IDs from an unexpanded menu are still not specification data.
    for identifier,sku in source.items():
        if identifier not in present and sku.get('skuSpec') and (len(source)>1 or data.get('haveMultiSpecs') in (1,'1')):
            groups.append({'skuId':sku.get('skuId'),'groupList':[]})
    for sku in groups:
        original=source.get(str(sku['skuId']),{})
        for name in ('skuSpec','originPrice','discountPrice','activityInfoList','discountPriceList','finalPrice','couponInfo','combinationSpec','combinationSpecList'):
            if name in original:sku[name]=deepcopy(original[name])
        by_group={str(g.get('groupId')):g for g in original.get('groupList') or []}
        for group in sku['groupList']:
            raw=by_group.get(str(group['groupId']),{})
            by_option={str(o.get('groupSkuId')):o for o in raw.get('groupSkuList') or []}
            for option in group['groupSkuList']:
                native=by_option.get(str(option['groupSkuId']),{})
                # Preserve nested structure and the numeric/currency price.
                for name in ('groupList','skuList','originPrice','minNumber','maxNumber','repeatable'):
                    if name in native:option[name]=deepcopy(native[name])
            if raw.get('groupList'):group['groupList']=deepcopy(raw['groupList'])
    order={identifier:i for i,identifier in enumerate(source)}
    groups.sort(key=lambda sku:order.get(str(sku['skuId']),len(order)))
    return groups


def unresolved_nested(data):
    if isinstance(data,list):return any(unresolved_nested(x) for x in data)
    if not isinstance(data,dict):return False
    if data.get('hasNestedGroup') not in (None,False,0,'0'):
        if not data.get('groupList') and not any(x.get('groupList') for x in data.get('groupSkuList') or []):return True
    return any(unresolved_nested(v) for v in data.values() if isinstance(v,(list,dict)))


def unavailable_products(menu):
    """Only explicit menu evidence; absent/unknown availability is not a skip."""
    products={str(p['spuId']):p for c in menu.get('shopCategoryList') or []
              for p in c.get('spuList') or [] if p.get('spuId') is not None}
    return {identifier for identifier,p in products.items()
            if p.get('name') and type(p.get('availableStatus')) in (int,str)
            and p['availableStatus'] in (0,'0')}


def coverage(menu,details):
    ids=set();products={}
    for cat in menu.get('shopCategoryList') or []:
        ids.update(str(x) for x in cat.get('spuIdList') or [])
        for product in cat.get('spuList') or []:
            identifier=str(product['spuId']);ids.add(identifier);products[identifier]=product
    missing_main=[];missing_details=[];nested=[]
    for identifier in sorted(ids):
        detail=details.get(identifier,{}).get('data') or {};product=products.get(identifier) or detail
        if not product.get('name'):missing_main.append(identifier)
        flag=product.get('haveMultiSpecs')
        if flag not in (0,'0') and not detail.get('name'):missing_details.append(identifier)
        if detail and unresolved_nested(detail):nested.append(identifier)
    return {'declared_products':len(ids),'missing_main_ids':missing_main,'missing_detail_ids':missing_details,
            'recorded_product_ids':sorted(ids-set(missing_main)),
            'unavailable_product_ids':sorted(unavailable_products(menu)),
            'nested_unresolved_ids':nested,'menu_complete':not missing_main,
            'custom_details_complete':not missing_details and not nested}


def apply_detail_policy(check,closed,unavailable_ids=()):
    # Keep raw missing/full-data coverage unchanged. Only the policy gate may
    # exempt named, recorded products, never missing menu rows or unknown IDs.
    exempt=set(unavailable_ids)&set(check.get('recorded_product_ids',[]))
    unresolved=set(check['missing_detail_ids'])|set(check['nested_unresolved_ids'])
    check['details_required']=not closed
    check['details_skip_reason']='store_closed' if closed else None
    check['unavailable_detail_ids']=sorted(unresolved&exempt) if not closed else []
    check['skipped_detail_ids']=sorted(unresolved if closed else unresolved&exempt)
    check['required_missing_detail_ids']=sorted(set(check['missing_detail_ids'])-exempt) if not closed else []
    check['required_nested_unresolved_ids']=sorted(set(check['nested_unresolved_ids'])-exempt) if not closed else []
    check['details_policy_complete']=bool(closed or not (unresolved-exempt))
    return check


def write_workbook(path,shop_rows,item_rows,coverage_rows,source_metadata):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    wb=openpyxl.Workbook();ws=wb.active;ws.title='店铺信息'
    overflow=[]
    for sheet,columns,rows in ((ws,CK.SHOP_COLS,shop_rows),(wb.create_sheet('菜品信息'),CK.ITEM_COLS,item_rows)):
        sheet.append(columns);sheet.freeze_panes='A2';sheet.auto_filter.ref=f'A1:{openpyxl.utils.get_column_letter(len(columns))}1'
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='334155')
        for row_number,row in enumerate(rows,2):
            values=[]
            for key in columns:
                value=row.get(key,'');value='' if value is None else value
                if isinstance(value,str) and len(value)>32767:
                    folder=path.with_suffix('.overflow');folder.mkdir(exist_ok=True)
                    filename=digest(value)+'.json';(folder/filename).write_text(value)
                    overflow.append({'shop_id':row.get('shop_id'),'item_id':row.get('item_id'),'column':key,'file':folder.name+'/'+filename})
                    value=compact({'full_value_file':folder.name+'/'+filename,'sha256':digest(value),'truncated':False})
                values.append(value)
            sheet.append(values)
            # Avoid max_row/max_column scanning all accumulated cells per row.
            for cell in next(sheet.iter_rows(min_row=row_number,max_row=row_number,min_col=1,max_col=len(columns))):
                if isinstance(cell.value,str):cell.data_type='s'
    wb.save(path)
    metadata=dict(source_metadata,shop_rows=len(shop_rows),item_rows=len(item_rows),shops=coverage_rows,overflow=overflow)
    path.with_suffix('.coverage.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2,default=str)+'\n')
    return {'path':str(path.resolve()),'shops':len(shop_rows),'items':len(item_rows),
            'complete_shop_jobs':sum(r['complete'] for r in coverage_rows),'partial_shop_jobs':sum(not r['complete'] for r in coverage_rows),'overflow_cells':len(overflow)}


def export_run(store,run_id,path):
    run=store.rows('SELECT settings FROM collection_runs WHERE id=%s',(run_id,))
    settings=unpack_json(run[0]['settings']) if run else {}
    closed=closed_shop_jobs(store,run_id) if settings.get('skip_closed_details') is True else set()
    jobs=store.rows('SELECT * FROM shop_jobs WHERE run_id=%s ORDER BY id',(run_id,))
    results=store.rows('SELECT r.*,a.user_id FROM task_results r JOIN shop_jobs j ON j.id=r.shop_job_id LEFT JOIN accounts a ON a.id=r.account_id WHERE j.run_id=%s ORDER BY r.observed_at,r.id',(run_id,))
    grouped={}
    for row in results:
        row['response']=json.loads(zlib.decompress(row['response_blob']));grouped.setdefault(row['shop_job_id'],[]).append(row)
    shops=[];items=[];checks=[]
    for job in jobs:
        raw=grouped.get(job['id'],[]);infos=[r for r in raw if r['endpoint']=='shopInfo'];menus=[r for r in raw if r['endpoint']=='productList']
        info=infos[-1] if infos else None;menu=menus[-1] if menus else None
        if info:
            row=CK.map_shop(info['response']['data'],job['shop_id'],job['latitude'],job['longitude'])
            row.update(user_type=1,account_name=info['user_id'] or '',create_time=info['observed_at'].isoformat(sep=' '),timestamp=int(info['observed_at'].replace(tzinfo=timezone.utc).timestamp()));shops.append(row)
        if menu:
            renders=[r['response'] for r in raw if r['endpoint']=='productRender' and r['observed_at']>=menu['observed_at']]
            merged=merge_menu(menu['response']['data'],renders)
            details={r['target_id']:r['response'] for r in raw if r['endpoint']=='productSpecifics'}
            check=coverage(merged,details)
            rows=CK.map_items(merged,job['shop_id'],details=details)
            for row in rows:
                detail=details.get(str(row['item_id']))
                if detail:
                    sub=sub_items(detail['data']);row['sub_item_json']=compact(sub) if sub else ''
                row.update(create_time=menu['observed_at'].isoformat(sep=' '),timestamp=int(menu['observed_at'].replace(tzinfo=timezone.utc).timestamp()))
            items.extend(rows)
        else:check={'menu_complete':False,'custom_details_complete':False,'declared_products':None,'missing_main_ids':[],'missing_detail_ids':[],'nested_unresolved_ids':[]}
        apply_detail_policy(check,job['id'] in closed)
        full=bool(info) and bool(menu) and check['menu_complete'] and check['custom_details_complete']
        check.update(shop_job_id=job['id'],shop_id=job['shop_id'],shop_info_complete=bool(info),
                     full_data_complete=full,
                     complete=bool(info) and bool(menu) and check['menu_complete'] and (check['custom_details_complete'] or not check['details_required']))
        checks.append(check)
    return write_workbook(path,shops,items,checks,{'run_id':run_id,'source':'mysql','timestamps':'UTC','atomic_cross_account_snapshot':False,'skip_closed_details':settings.get('skip_closed_details') is True})


def export_status(store,path):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    tables=[('账号状态',store.rows('SELECT * FROM account_dashboard ORDER BY id')),
            ('每日接口统计',store.rows('SELECT * FROM daily_account_requests ORDER BY business_date,account_id,endpoint,origin')),
            ('每日预算',store.rows('SELECT a.label,p.endpoint,p.work_limit,p.hard_limit,p.timezone FROM budget_policies p JOIN accounts a ON a.id=p.account_id ORDER BY a.id,p.endpoint')),
            ('队列状态',store.rows('SELECT j.run_id,t.endpoint,t.state,COUNT(*) AS task_count FROM tasks t JOIN shop_jobs j ON j.id=t.shop_job_id GROUP BY j.run_id,t.endpoint,t.state'))]
    wb=openpyxl.Workbook();wb.remove(wb.active)
    for name,rows in tables:
        ws=wb.create_sheet(name)
        if not rows:continue
        keys=list(rows[0]);ws.append(keys);ws.freeze_panes='A2'
        for row in rows:
            ws.append([row[k] for k in keys])
    wb.save(path)
    return {'path':str(path.resolve()),'sheets':[x[0] for x in tables]}


def cost_job_scope(store,run_id):
    """Follow reused data to original requests, retaining the selected shops."""
    root=store.rows('SELECT shop_id,context_key FROM shop_jobs WHERE run_id=%s',(run_id,))
    wanted={(r['shop_id'],r['context_key']) for r in root};jobs=set();seen=set();lineage=[]
    current=run_id
    while current is not None and current not in seen:
        seen.add(current)
        rows=store.rows('SELECT settings FROM collection_runs WHERE id=%s',(current,))
        if not rows:break
        scoped=store.rows('SELECT id,shop_id,context_key FROM shop_jobs WHERE run_id=%s',(current,))
        selected=[r['id'] for r in scoped if (r['shop_id'],r['context_key']) in wanted]
        jobs.update(selected);lineage.append({'run_id':current,'shop_jobs':len(selected)})
        current=unpack_json(rows[0]['settings']).get('source_run_id')
    return sorted(jobs),lineage


def summarize_request_costs(attempts,refreshes,rates,selected_accounts,run_id,delivery=None):
    from collections import Counter,defaultdict
    import math
    rates={k:rates.get(k) for k in ('account_unit_price','traffic_price_per_gb','currency')}
    for key in ('account_unit_price','traffic_price_per_gb'):
        if rates[key] is not None:
            rates[key]=float(rates[key])
            if not math.isfinite(rates[key]) or rates[key]<0:raise ValueError('invalid unit price')
    attempts=list({r['id']:r for r in attempts}.values())
    endpoints={};accounts={};samples=defaultdict(list);missing=Counter()
    measured=up=down=0;provider_bytes=Counter();estimated_by_provider=Counter()
    for row in attempts:
        endpoint=row['endpoint'];aid=row['account_id'];outcome=row['outcome'];provider=row.get('cost_provider') or 'unknown'
        base={'requests':0,'current_requests':0,'successful_requests':0,'rejected_requests':0,'closed_requests':0,'transport_errors':0,'measured_requests':0,'uploaded_bytes':0,'downloaded_bytes':0}
        ep=endpoints.setdefault(endpoint,dict(base,endpoint=endpoint))
        ac=accounts.setdefault(aid,dict(base,account_id=aid))
        available=row.get('uploaded_bytes') is not None and row.get('downloaded_bytes') is not None
        for group in (ep,ac):
            group['requests']+=1;group['current_requests']+=int(row['run_id']==run_id)
            group['successful_requests']+=int(outcome=='success');group['rejected_requests']+=int(outcome=='rejected')
            group['closed_requests']+=int(outcome=='store_closed');group['transport_errors']+=int(outcome=='transport_error')
            if available:
                group['measured_requests']+=1;group['uploaded_bytes']+=int(row['uploaded_bytes']);group['downloaded_bytes']+=int(row['downloaded_bytes'])
        if available:
            measured+=1;up+=int(row['uploaded_bytes']);down+=int(row['downloaded_bytes'])
            size=int(row['uploaded_bytes'])+int(row['downloaded_bytes'])
            provider_bytes[provider]+=size;samples[endpoint,outcome,provider].append(size)
        else:missing[endpoint,outcome,provider]+=1
    extra_up=sum(int(r['uploaded_bytes'] or 0) for r in refreshes)
    extra_down=sum(int(r['downloaded_bytes'] or 0) for r in refreshes)
    for r in refreshes:provider_bytes[r.get('cost_provider') or 'unknown']+=int(r['uploaded_bytes'] or 0)+int(r['downloaded_bytes'] or 0)
    estimated_missing=0;unestimated=0
    for key,count in missing.items():
        if samples[key]:
            estimate=sum(samples[key])/len(samples[key])*count
            estimated_missing+=estimate;estimated_by_provider[key[2]]+=estimate
        else:unestimated+=count
    total=len(attempts);used=len(accounts);selected=len(set(selected_accounts));gb=(up+down+extra_up+extra_down)/1e9
    cost=lambda quantity,price:None if price is None else round(quantity*price,8)
    return {'run_id':run_id,'rates':rates,'selected_accounts':selected,'used_accounts':used,
        'total_requests_including_reused_data':total,
        'current_run_requests':sum(r['current_requests'] for r in endpoints.values()),
        'successful_requests':sum(r['successful_requests'] for r in endpoints.values()),
        'measured_requests':measured,'unmeasured_requests':total-measured,
        'uploaded_bytes':up+extra_up,'downloaded_bytes':down+extra_down,'measured_proxy_stream_gb':gb,
        'refresh_count':len(refreshes),'unmeasured_refreshes':sum(r.get('uploaded_bytes') is None for r in refreshes),
        'estimated_proxy_stream_gb':(up+down+extra_up+extra_down+estimated_missing)/1e9,
        'unestimated_requests':unestimated,
        'touched_account_purchase_cost':cost(used,rates['account_unit_price']),
        'selected_pool_purchase_cost':cost(selected,rates['account_unit_price']),
        'measured_traffic_cost':cost(provider_bytes['ipfoxy']/1e9,rates['traffic_price_per_gb']),
        'ipfoxy_measured_gb':provider_bytes['ipfoxy']/1e9,
        'clash_subscription_measured_gb':provider_bytes['clash_subscription']/1e9,
        'unknown_provider_measured_gb':provider_bytes['unknown']/1e9,
        'unpriced_traffic_gb':(provider_bytes['clash_subscription']+provider_bytes['unknown'])/1e9,
        'ipfoxy_estimated_gb':(provider_bytes['ipfoxy']+estimated_by_provider['ipfoxy'])/1e9,
        'complete_shops':(delivery or {}).get('complete_shop_jobs'),'items':(delivery or {}).get('items'),
        'endpoints':sorted(endpoints.values(),key=lambda x:x['endpoint']),
        'accounts':sorted(accounts.values(),key=lambda x:x['account_id']),
        'measurement_scope':'local_proxy_stream',
        'notes':['Measured TCP stream bytes at the local proxy boundary include CONNECT/TLS data, but exclude IP/TCP packet headers and retransmissions.',
                 'These are collection measurements, not IPFoxy billing counters. Pricing assumes combined upload/download and decimal GB (1e9 bytes).',
                 'The configured per-GB rate applies only to identified IPFoxy traffic. Clash subscription and unknown-provider traffic remain unpriced, not free.',
                 'Historical traffic remains unmeasured. Estimates use only matching endpoint/outcome samples; unmatched requests remain unestimated.',
                 'Used accounts means accounts with sent requests, not accounts permanently exhausted. Account acquisition/login traffic is outside this measurement.',
                 'Reused data includes its original request history once; current_run_requests counts new requests in this batch.']}


def export_costs(store,run_id,folder,delivery=None):
    jobs,lineage=cost_job_scope(store,run_id);attempts=[];refreshes=[]
    if jobs:
        placeholders=','.join(['%s']*len(jobs))
        attempts=store.rows(f"SELECT a.id,a.account_id,a.endpoint,a.outcome,j.run_id,m.uploaded_bytes,m.downloaded_bytes,m.cost_provider FROM request_attempts a JOIN tasks t ON t.id=a.task_id JOIN shop_jobs j ON j.id=t.shop_job_id LEFT JOIN network_measurements m ON m.attempt_id=a.id WHERE j.id IN ({placeholders}) AND a.counts_budget=TRUE AND a.state IN ('sent','done','uncertain')",jobs)
        refreshes=store.rows(f"SELECT e.id,m.uploaded_bytes,m.downloaded_bytes,m.cost_provider FROM proxy_refresh_events e JOIN tasks t ON t.id=e.task_id LEFT JOIN network_measurements m ON m.proxy_refresh_id=e.id WHERE t.shop_job_id IN ({placeholders})",jobs)
    selected=store.rows('SELECT account_ids FROM executions WHERE run_id=%s ORDER BY id DESC LIMIT 1',(run_id,))
    ids=unpack_json(selected[0]['account_ids']) if selected else []
    report=summarize_request_costs(attempts,refreshes,store.get_setting('cost_rates') or {},ids,run_id,delivery)
    report['lineage']=lineage
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    (folder/'costs.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    wb=openpyxl.Workbook();ws=wb.active;ws.title='成本概览'
    rows=[
        ['指标','数值','说明'],
        ['单账号采购价',report['rates']['account_unit_price'],'可填写实际采购价'],
        ['IPFoxy每GB价格',report['rates']['traffic_price_per_gb'],'上下行合计，1GB=10^9字节；不用于Clash订阅'],
        ['币种',report['rates']['currency'] or '待填写',''],
        ['选定账号池数量',report['selected_accounts'],'已选账号，未必均已使用'],
        ['实际触达账号数',report['used_accounts'],'不等于账号已耗尽'],
        ['触达账号采购成本','=IF(ISNUMBER(B2),B2*B6,"")',''],
        ['整个账号池采购成本','=IF(ISNUMBER(B2),B2*B5,"")',''],
        ['总请求数（含复用来源）',report['total_requests_including_reused_data'],'来源请求去重计入'],
        ['本批新增请求',report['current_run_requests'],''],
        ['继承数据对应请求',report['total_requests_including_reused_data']-report['current_run_requests'],''],
        ['已计量业务请求',report['measured_requests'],''],
        ['未计量历史请求',report['unmeasured_requests'],'不是零流量'],
        ['已测上行字节',report['uploaded_bytes'],'含代理刷新请求'],
        ['已测下行字节',report['downloaded_bytes'],''],
        ['已测代理连接流量GB',report['measured_proxy_stream_gb'],'本地代理流字节，不是供应商账单'],
        ['IPFoxy已测流量成本','=IF(ISNUMBER(B3),B3*B24,"")','其他出口费用待定'],
        ['IPFoxy估算流量GB',report['ipfoxy_estimated_gb'],'仍有 '+str(report['unestimated_requests'])+' 次历史请求未估算'],
        ['估算流量成本','=IF(ISNUMBER(B3),B3*B18,"")',''],
        ['按本轮规则完整店铺',report['complete_shops'],'未完成时只能作阶段样本'],
        ['已导出菜品数',report['items'],''],
        ['触达账号+已定价流量','=IF(AND(ISNUMBER(B2),ISNUMBER(B3)),B7+B17,"")','不含Clash订阅、未知出口、未计量部分、服务器和人工费用'],
        ['每完整店阶段成本','=IF(AND(ISNUMBER(B22),B20>0),B22/B20,"")','仅已定价部分，整批完成后再用于结算'],
        ['IPFoxy已测流量GB',report['ipfoxy_measured_gb'],''],
        ['Clash订阅已测流量GB',report['clash_subscription_measured_gb'],'订阅成本未提供，不按45元/GB计费'],
        ['未知出口已测流量GB',report['unknown_provider_measured_gb'],'来源无法确认时不猜测供应商'],
        ['未定价流量GB',report['unpriced_traffic_gb'],'不是免费流量']]
    for row in rows:ws.append(row)
    ws['B4'].data_type='s'
    for name,key in (('接口统计','endpoints'),('账号统计','accounts')):
        sheet=wb.create_sheet(name);data=report[key]
        if data:
            keys=list(data[0]);sheet.append(keys)
            for row in data:sheet.append([row[k] for k in keys])
    for sheet in wb:
        sheet.freeze_panes='A2'
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='334155')
        sheet.column_dimensions['A'].width=32;sheet.column_dimensions['B'].width=22;sheet.column_dimensions['C'].width=72
    wb.save(folder/'costs.xlsx')
    return {k:report[k] for k in ('selected_accounts','used_accounts','measured_requests','unmeasured_requests','measured_proxy_stream_gb','measured_traffic_cost')}
